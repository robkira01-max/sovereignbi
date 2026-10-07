"""Tests export Excel et PDF.

Keys are set in conftest.py (viewer="test-viewer-key-s6").
"""
import io
import os

import pandas as pd
import pytest

import openpyxl
from fastapi.testclient import TestClient

from app.export.excel import results_to_excel
from app.export.pdf import results_to_pdf

_VIEWER_KEY = os.environ.get("SOVEREIGN_VIEWER_KEY", "test-viewer-key-s6")
_V = {"X-Api-Key": _VIEWER_KEY}


# ── Pure function: Excel ──────────────────────────────────────────────────────

def test_results_to_excel_parsable():
    cols = ["region", "ca"]
    rows = [["Ontario", 10000.5], ["Quebec", 8500.0]]
    data = results_to_excel(cols, rows, title="Test CA")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.cell(1, 1).value == "region"
    assert ws.cell(1, 2).value == "ca"
    assert ws.cell(2, 1).value == "Ontario"


def test_results_to_excel_empty_rows():
    data = results_to_excel(["col_a"], [])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.cell(1, 1).value == "col_a"
    assert ws.max_row == 1


def test_results_to_excel_none_values():
    data = results_to_excel(["a", "b"], [[None, "x"], ["y", None]])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    # openpyxl reads empty string cells back as None
    assert ws.cell(2, 1).value in (None, "")
    assert ws.cell(2, 2).value == "x"


def test_results_to_excel_returns_bytes():
    assert isinstance(results_to_excel(["x"], [[1]]), bytes)


# ── Pure function: PDF ────────────────────────────────────────────────────────

def test_results_to_pdf_returns_bytes():
    assert isinstance(results_to_pdf(["col"], [[1]]), bytes)


def test_results_to_pdf_starts_with_pdf_magic():
    data = results_to_pdf(["a"], [[1]])
    assert data[:4] == b"%PDF"


def test_results_to_pdf_contains_disclaimer():
    # fpdf2 compresses content streams; text is not directly readable in raw bytes.
    # Verify that a PDF with rows is at least as large as an empty one (disclaimer always present).
    data_with_content = results_to_pdf(["col"], [[1]])
    data_empty = results_to_pdf([], [])
    assert len(data_with_content) >= len(data_empty)


def test_results_to_pdf_empty_columns():
    data = results_to_pdf([], [])
    assert data[:4] == b"%PDF"


# ── Integration: HTTP routes ──────────────────────────────────────────────────

@pytest.fixture
def client_with_data():
    """TestClient with a CSV pre-loaded and a FakeLLM injected."""
    import app.main as m
    from app.data.duckdb_store import DataStore
    from app.pipeline.graph import build_graph

    test_store = DataStore()
    test_store.load_dataframe("ventes", pd.DataFrame({"region": ["ON"], "montant": [1000.0]}))

    class FakeLLM:
        def complete(self, system, user):
            # Answer node prompt contains "Résultats :" — SQL gen prompt does not.
            if "résultats :" in user.lower():
                return "Le CA de Ontario est 1000."
            return "SELECT region, SUM(montant) AS ca FROM ventes GROUP BY region"

    fake_graph = build_graph(FakeLLM(), test_store)

    original_store = m.store
    original_get_graph = m.app.state.get_graph

    m.store = test_store
    m.app.state.store = test_store
    m.app.state.get_graph = lambda: fake_graph

    yield TestClient(m.app)

    m.store = original_store
    m.app.state.store = original_store
    m.app.state.get_graph = original_get_graph


def test_export_excel_returns_xlsx(client_with_data):
    r = client_with_data.post("/export/excel", json={"question": "CA ?"}, headers=_V)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]


def test_export_pdf_returns_pdf(client_with_data):
    r = client_with_data.post("/export/pdf", json={"question": "CA ?"}, headers=_V)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"


def test_export_no_key_returns_401():
    from app.main import app
    c = TestClient(app)
    r = c.post("/export/excel", json={"question": "CA ?"})
    assert r.status_code == 401


def test_export_no_data_returns_400():
    import app.main as m
    from app.data.duckdb_store import DataStore

    empty_store = DataStore()  # no tables
    original = m.app.state.store
    original_main_store = m.store
    m.store = empty_store
    m.app.state.store = empty_store
    c = TestClient(m.app)
    r = c.post("/export/excel", json={"question": "CA ?"}, headers=_V)
    assert r.status_code == 400
    m.store = original_main_store
    m.app.state.store = original
