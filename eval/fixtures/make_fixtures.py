"""Generate all synthetic fixture files for Sprint 2 evaluation.

Run: python eval/fixtures/make_fixtures.py
All data is purely synthetic — no real customer data.
"""
import hashlib
import random
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font

HERE = Path(__file__).parent
RNG = np.random.default_rng(42)
random.seed(42)

# ── Constants ────────────────────────────────────────────────────────────────
REGIONS = ["Ontario", "Québec", "BC", "Alberta", "Manitoba"]
CANAUX = ["en_ligne", "magasin", "telephone"]
CATEGORIES = ["Mobilier", "Électronique", "Vêtements", "Alimentation", "Sport"]
MODES_PAIEMENT = ["carte", "virement", "cheque", "especes"]
STATUTS = ["confirme", "retourne", "annule", "en_cours"]
TRANSPORTEURS = ["Purolator", "Canada Post", "FedEx", "UPS", "DHL"]
PROVINCES = ["Ontario", "Québec", "BC", "Alberta", "Manitoba", "Saskatchewan", "Nouvelle-Écosse"]
DEPARTEMENTS = ["Ventes", "Marketing", "Finance", "IT", "RH", "Logistique", "Direction"]
SOURCES_LEAD = ["google", "referral", "direct", "email", "social"]
CAMPAGNES = ["printemps2026", "noel2025", "soldes", "fidélité", "aucune"]
SEGMENTS = ["premium", "standard", "prospect", "inactif"]


def _fake_id(prefix: str, n: int) -> list[str]:
    return [f"{prefix}{i:05d}" for i in range(1, n + 1)]


# ── Clients ──────────────────────────────────────────────────────────────────
def make_clients(n: int = 200) -> pd.DataFrame:
    noms = ["Tremblay", "Gagnon", "Roy", "Côté", "Bouchard", "Smith", "Johnson",
            "Williams", "Brown", "Jones", "Martin", "Bernard", "Dubois", "Thomas"]
    prenoms = ["Marie", "Jean", "Pierre", "Sophie", "Michel", "Julie", "François",
               "Isabelle", "Luc", "Emma", "Noah", "Olivia", "Liam", "Ava"]
    ids = _fake_id("CLI", n)
    df = pd.DataFrame({
        "id_client": ids,
        "nom": RNG.choice(noms, n),
        "prenom": RNG.choice(prenoms, n),
        "email": [f"{ids[i].lower()}@example.com" for i in range(n)],
        "province": RNG.choice(PROVINCES, n),
        "code_postal": [f"{chr(RNG.integers(65,90))}0{RNG.integers(1,9)} {RNG.integers(1,9)}{chr(RNG.integers(65,90))}{RNG.integers(1,9)}"
                        for _ in range(n)],
        "type_client": RNG.choice(["entreprise", "particulier"], n, p=[0.3, 0.7]),
        "age": RNG.integers(18, 75, n),
        "genre": RNG.choice(["M", "F", "NB"], n, p=[0.48, 0.48, 0.04]),
        "langue": RNG.choice(["fr", "en"], n, p=[0.4, 0.6]),
        "date_creation": pd.to_datetime(
            RNG.integers(pd.Timestamp("2020-01-01").value,
                         pd.Timestamp("2026-10-01").value, n)),
        "actif": RNG.choice([True, False], n, p=[0.9, 0.1]),
    })
    return df


# ── Produits ─────────────────────────────────────────────────────────────────
def make_produits(n: int = 80) -> pd.DataFrame:
    noms_par_cat = {
        "Mobilier": ["Bureau", "Chaise", "Étagère", "Table", "Armoire"],
        "Électronique": ["Laptop", "Tablette", "Écran", "Clavier", "Souris"],
        "Vêtements": ["Veste", "Pantalon", "T-shirt", "Robe", "Manteau"],
        "Alimentation": ["Café", "Thé", "Huile", "Farine", "Sucre"],
        "Sport": ["Vélo", "Tapis", "Haltère", "Balle", "Raquette"],
    }
    rows = []
    for i in range(1, n + 1):
        cat = RNG.choice(CATEGORIES)
        nom = RNG.choice(noms_par_cat[cat]) + f" {i}"
        prix = round(float(RNG.uniform(10, 1500)), 2)
        rows.append({
            "code_produit": f"PRD{i:04d}",
            "nom_produit": nom,
            "categorie": cat,
            "prix_catalogue": prix,
            "cout_unitaire": round(prix * float(RNG.uniform(0.3, 0.6)), 2),
            "poids_kg": round(float(RNG.uniform(0.1, 50)), 3),
            "quantite_stock": int(RNG.integers(0, 500)),
            "id_fournisseur": f"FOUR{RNG.integers(1, 15):03d}",
            "actif": bool(RNG.choice([True, False], p=[0.85, 0.15])),
            "unite": RNG.choice(["pièce", "kg", "litre", "boîte"]),
        })
    return pd.DataFrame(rows)


# ── Commandes ─────────────────────────────────────────────────────────────────
def make_commandes(client_ids: list[str], n: int = 300) -> pd.DataFrame:
    ids = _fake_id("CMD", n)
    date_cmd = pd.to_datetime(
        RNG.integers(pd.Timestamp("2025-01-01").value,
                     pd.Timestamp("2026-10-01").value, n))
    delais = RNG.integers(1, 15, n)
    date_prev = date_cmd + pd.to_timedelta(delais, unit="D")
    retard = RNG.integers(-3, 8, n)
    date_reelle = date_prev + pd.to_timedelta(retard, unit="D")
    statut_liv = RNG.choice(["livre", "en_transit", "retourne", "annule"], n, p=[0.7, 0.15, 0.1, 0.05])
    return pd.DataFrame({
        "id_commande": ids,
        "id_client": RNG.choice(client_ids, n),
        "date_commande": date_cmd.date,
        "date_livraison_prevue": date_prev.date,
        "date_livraison_reelle": date_reelle.date,
        "statut_livraison": statut_liv,
        "transporteur": RNG.choice(TRANSPORTEURS, n),
        "cout_livraison": np.round(RNG.uniform(5, 80, n), 2),
        "delai_livraison_jours": delais,
    })


# ── Employés ──────────────────────────────────────────────────────────────────
def make_employes(n: int = 50) -> pd.DataFrame:
    ids = _fake_id("EMP", n)
    noms = ["Côté", "Lavoie", "Pelletier", "Morin", "Lefebvre",
            "Miller", "Wilson", "Moore", "Taylor", "Anderson"]
    return pd.DataFrame({
        "id_employe": ids,
        "nom": [RNG.choice(noms) + f" {i}" for i in range(n)],
        "departement": RNG.choice(DEPARTEMENTS, n),
        "date_embauche": pd.to_datetime(
            RNG.integers(pd.Timestamp("2010-01-01").value,
                         pd.Timestamp("2026-01-01").value, n)).date,
        "salaire_annuel": np.round(RNG.uniform(42000, 120000, n), 2),
        "type_contrat": RNG.choice(["CDI", "CDD", "independant"], n, p=[0.7, 0.2, 0.1]),
        "titre": RNG.choice(["Analyste", "Responsable", "Directeur", "Coordinateur", "Spécialiste"], n),
        "manager_id": [ids[int(RNG.integers(0, n))] if RNG.random() > 0.1 else None for _ in range(n)],
    })


# ── Ventes détail (schéma large) ──────────────────────────────────────────────
def make_ventes_detail(client_ids: list[str], produit_ids: list[str],
                       commande_ids: list[str], employe_ids: list[str],
                       n: int = 500) -> pd.DataFrame:
    quantites = RNG.integers(1, 20, n)
    prix_ht = np.round(RNG.uniform(10, 1500, n), 2)
    taux_remise = np.round(RNG.choice([0.0, 0.05, 0.10, 0.15, 0.20, 0.25, 0.30], n,
                                      p=[0.4, 0.15, 0.15, 0.1, 0.1, 0.05, 0.05]), 2)
    remise_montant = np.round(quantites * prix_ht * taux_remise, 2)
    montant_ht = np.round(quantites * prix_ht * (1 - taux_remise), 2)
    tva = 0.15
    montant_ttc = np.round(montant_ht * (1 + tva), 2)
    marge = np.round(montant_ht * RNG.uniform(0.1, 0.5, n), 2)

    date_vente = pd.to_datetime(
        RNG.integers(pd.Timestamp("2025-01-01").value,
                     pd.Timestamp("2026-10-01").value, n))

    df = pd.DataFrame({
        # Core
        "id_ligne": range(1, n + 1),
        "id_commande": RNG.choice(commande_ids, n),
        "id_client": RNG.choice(client_ids, n),
        "id_vendeur": RNG.choice(employe_ids, n),
        "date_vente": date_vente.date,
        "region": RNG.choice(REGIONS, n),
        "canal": RNG.choice(CANAUX, n),
        "code_produit": RNG.choice(produit_ids, n),
        "nom_produit": [f"Produit {p}" for p in RNG.choice(produit_ids, n)],
        "categorie_produit": RNG.choice(CATEGORIES, n),
        "quantite": quantites,
        "prix_unitaire_ht": prix_ht,
        "taux_remise": taux_remise,
        "remise_montant": remise_montant,
        "montant_ht": montant_ht,
        "tva": tva,
        "montant_ttc": montant_ttc,
        "mode_paiement": RNG.choice(MODES_PAIEMENT, n),
        "statut": RNG.choice(STATUTS, n, p=[0.75, 0.1, 0.05, 0.1]),
        # Extended (30+ colonnes supplémentaires)
        "code_postal_livraison": [f"A{RNG.integers(1,9)}B {RNG.integers(1,9)}C{RNG.integers(1,9)}" for _ in range(n)],
        "pays_livraison": "Canada",
        "poids_total_kg": np.round(RNG.uniform(0.1, 100, n), 2),
        "volume_cm3": np.round(RNG.uniform(100, 50000, n), 0),
        "transporteur": RNG.choice(TRANSPORTEURS, n),
        "numero_suivi": [f"TRK{RNG.integers(1000000, 9999999)}" for _ in range(n)],
        "date_expedition": (date_vente + pd.to_timedelta(RNG.integers(0, 3, n), unit="D")).date,
        "delai_livraison_jours": RNG.integers(1, 14, n),
        "notes_commande": RNG.choice(["", "Urgent", "Fragile", "Cadeau", ""], n),
        "priorite": RNG.choice(["normale", "haute", "basse"], n),
        "source_lead": RNG.choice(SOURCES_LEAD, n),
        "campagne_marketing": RNG.choice(CAMPAGNES, n),
        "code_promo": RNG.choice(["", "PROMO10", "ETE25", "NOEL15", ""], n, p=[0.6, 0.1, 0.1, 0.1, 0.1]),
        "taux_satisfaction": np.where(RNG.random(n) > 0.2, np.round(RNG.uniform(1, 5, n), 1), np.nan),
        "note_client": RNG.choice(["", "Très satisfait", "Satisfait", "Neutre", "Insatisfait"], n),
        "nb_articles_retournes": np.where(RNG.random(n) > 0.9, RNG.integers(1, 5, n), 0),
        "motif_retour": RNG.choice(["", "Défaut", "Erreur commande", "Non conforme", ""], n),
        "id_contrat": RNG.choice([f"CTR{i:04d}" for i in range(1, 21)] + [None] * 10, n),
        "type_contrat_client": RNG.choice(["spot", "cadre", "abonnement"], n),
        "devise": "CAD",
        "taux_change": 1.0,
        "montant_devise_etrangere": montant_ttc,
        "marge_brute": marge,
        "taux_marge": np.round(marge / np.where(montant_ht > 0, montant_ht, 1), 4),
        "cout_acquisition_client": np.round(RNG.uniform(0, 500, n), 2),
        "valeur_vie_client": np.round(RNG.uniform(100, 10000, n), 2),
        "segment_client": RNG.choice(SEGMENTS, n),
        "score_risque": np.round(RNG.uniform(0, 1, n), 3),
        "indicateur_fraude": np.where(RNG.random(n) > 0.99, 1, 0),
        "hash_transaction": [hashlib.sha256(f"tx{i}".encode()).hexdigest()[:16] for i in range(n)],
        "langue_interface": RNG.choice(["fr", "en"], n),
    })
    assert len(df.columns) >= 50, f"Schéma insuffisant : {len(df.columns)} colonnes"
    return df


# ── Excel fixtures ────────────────────────────────────────────────────────────
def make_excel_simple(path: Path) -> None:
    df = pd.DataFrame({
        "produit": ["Bureau", "Chaise", "Table"],
        "quantite": [10, 25, 5],
        "prix": [299.99, 149.99, 399.99],
    })
    df.to_excel(path, index=False)


def make_excel_multi_header(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    # Row 1: category headers
    ws["A1"] = "Produit"
    ws["B1"] = "Ventes"
    ws["C1"] = "Ventes"
    ws["D1"] = "Ventes"
    ws["B1"].font = Font(bold=True)
    # Row 2: sub-headers
    ws["A2"] = "Nom"
    ws["B2"] = "Q1"
    ws["C2"] = "Q2"
    ws["D2"] = "Q3"
    # Data
    data = [("Bureau", 100, 120, 90), ("Chaise", 200, 180, 220), ("Table", 50, 60, 45)]
    for row in data:
        ws.append(row)
    wb.save(path)


def make_excel_multi_sheet(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Ventes"
    ws1.append(["date", "montant", "region"])
    ws1.append(["2026-01-01", 1000, "Ontario"])
    ws1.append(["2026-01-02", 2000, "Québec"])

    ws2 = wb.create_sheet("Clients")
    ws2.append(["id_client", "nom", "province"])
    ws2.append(["CLI001", "Tremblay", "Québec"])

    ws3 = wb.create_sheet("Résumé")
    ws3.append(["metric", "valeur"])
    ws3.append(["CA total", 3000])
    wb.save(path)


def make_excel_leading_empty_rows(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    # 5 blank rows, then header+data
    for _ in range(5):
        ws.append([None, None, None])
    ws.append(["produit", "quantite", "montant"])
    ws.append(["Bureau", 10, 2999.90])
    ws.append(["Chaise", 25, 3749.75])
    wb.save(path)


def make_excel_wide_schema(path: Path, n_cols: int = 55) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = [f"col_{i:02d}" for i in range(n_cols)]
    ws.append(headers)
    for row_i in range(10):
        ws.append([row_i * n_cols + c for c in range(n_cols)])
    wb.save(path)


# ── Main ──────────────────────────────────────────────────────────────────────
def main() -> None:
    print("Generating synthetic fixtures…")

    clients_df = make_clients(200)
    clients_df.to_csv(HERE / "clients.csv", index=False)
    print(f"  clients.csv — {len(clients_df)} rows")

    produits_df = make_produits(80)
    produits_df.to_csv(HERE / "produits.csv", index=False)
    print(f"  produits.csv — {len(produits_df)} rows")

    commandes_df = make_commandes(clients_df["id_client"].tolist(), 300)
    commandes_df.to_csv(HERE / "commandes.csv", index=False)
    print(f"  commandes.csv — {len(commandes_df)} rows")

    employes_df = make_employes(50)
    employes_df.to_csv(HERE / "employes.csv", index=False)
    print(f"  employes.csv — {len(employes_df)} rows")

    ventes_df = make_ventes_detail(
        clients_df["id_client"].tolist(),
        produits_df["code_produit"].tolist(),
        commandes_df["id_commande"].tolist(),
        employes_df["id_employe"].tolist(),
        500,
    )
    ventes_df.to_csv(HERE / "large_schema.csv", index=False)
    print(f"  large_schema.csv — {len(ventes_df)} rows × {len(ventes_df.columns)} cols")

    # Excel fixtures
    make_excel_simple(HERE / "classeur_simple.xlsx")
    print("  classeur_simple.xlsx")
    make_excel_multi_header(HERE / "classeur_en_tetes_multi.xlsx")
    print("  classeur_en_tetes_multi.xlsx")
    make_excel_multi_sheet(HERE / "classeur_multi_feuilles.xlsx")
    print("  classeur_multi_feuilles.xlsx")
    make_excel_leading_empty_rows(HERE / "classeur_lignes_vides.xlsx")
    print("  classeur_lignes_vides.xlsx")
    make_excel_wide_schema(HERE / "classeur_colonnes_larges.xlsx")
    print("  classeur_colonnes_larges.xlsx")

    print("Done.")


if __name__ == "__main__":
    main()
